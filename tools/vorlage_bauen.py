#!/usr/bin/env python3
"""Baut aus einer Vorlagen-Beschreibung die Download-Dateien mit Logo und Wasserzeichen.

Der Vorlagen-Agent schreibt den Inhalt einer Checkliste, eines Protokolls,
eines Musterschreibens oder einer Tabelle als YAML-Datei nach
``vorlagen/<kurzform>.yml``. Dieses Skript macht daraus die Dateien, die
Leser herunterladen:

    PDF zum Ausdrucken   – Kästchen und Linien zum Ankreuzen von Hand
    PDF ausfüllbar       – echte Formularfelder (Ankreuzen und Tippen am Bildschirm)
    Word (.docx)         – zum Anpassen, mit Logo im Kopf und Wasserzeichen
    Excel (.xlsx)        – für Tabellen, mit Logo und Fußzeile

Jede Datei trägt das Logo des Büros, den Stand und den Hinweis auf
ing-bassam.de; PDF und Word zusätzlich ein dezentes Wasserzeichen. Schriften
und Farben entsprechen der Website (tools/schriften/, fachwissen/artikel.css).

Aufruf:
    python tools/vorlage_bauen.py vorlagen/<kurzform>.yml [--ziel vorlagen/<kurzform>]
    python tools/vorlage_bauen.py --pruefen vorlagen/<kurzform>.yml   (nur Beschreibung prüfen)

Ausgabe: die erzeugten Dateien, eine je Zeile. Rückgabewert 2 bei einer
fehlerhaften Beschreibung. Benötigt reportlab, python-docx, openpyxl, Pillow
und PyYAML – alle kostenlos.

Aufbau der Beschreibung (YAML):

    titel: Checkliste Endabnahme Haus
    untertitel: Abnahmeprotokoll für Bauherren          # optional
    kurzform: abnahmeprotokoll-hausbau
    typ: Checkliste            # Checkliste | Protokoll | Musterschreiben | Tabelle
    stand: 2026-09-24
    dateien: [pdf, pdf-ausfuellbar, docx]              # gewünschte Formate; Tabelle: xlsx
    einleitung: Ein bis drei Sätze, was die Vorlage leistet und wie man sie nutzt.
    kopffelder: [Objekt und Anschrift, Datum, Anwesende]   # Eingabefelder oben (nicht bei Musterschreiben)
    abschnitte:                # Checkliste, Protokoll
      - titel: Außenanlagen
        punkte:
          - text: Regenwasser läuft vom Gebäude weg
            hinweis: Gefälle vom Haus weg, keine Pfützen an der Fassade   # optional
    unterschriften: [Bauherr, Unternehmer]             # optional, Protokoll
    tabelle:                   # Tabelle (auch als Zusatz zu anderen Typen)
      spalten: [Nr., Bauteil, Intervall, Zuständig, Erledigt am, Bemerkung]
      breiten: [1, 4, 2, 2, 2, 4]                       # relative Breiten, optional
      zeilen:                                            # vorgegebene Zeilen, optional
        - ["1", "Dachrinnen und Fallrohre", "jährlich", "", "", ""]
      leerzeilen: 15                                     # zusätzliche leere Zeilen
    brief:                     # Musterschreiben
      betreff: Bedenkenanzeige nach § 4 Abs. 3 VOB/B – Bauvorhaben [Bezeichnung]
      anrede: Sehr geehrte Damen und Herren,
      absaetze:
        - "hiermit melden wir Bedenken gegen … [Beschreibung] …"
      gruss: Mit freundlichen Grüßen
      anlagen: [Fotos, Planauszug]                       # optional
    hinweise:                  # Kurze Erläuterungen am Ende, optional
      - "Platzhalter in eckigen Klammern durch eigene Angaben ersetzen."
    quellen:                   # optional, klein am Ende
      - "§ 4 Abs. 3 VOB/B, Ausgabe 2016"

YAML-Regel: Jeder Text, der eckige Klammern, einen Doppelpunkt mit Leerzeichen
oder ein führendes Sonderzeichen enthält, steht in doppelten Anführungszeichen –
sonst liest YAML „[Datum]“ als Liste.
"""
from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path

import yaml

WURZEL = Path(__file__).resolve().parent.parent
SCHRIFTEN = WURZEL / "tools" / "schriften"

FIRMA_KURZ = "BIB Ingenieurbüro für Bauwesen"
FIRMA_LANG = "Bassam Ingenieurbüro für Bauwesen GmbH"
WEBSITE = "ing-bassam.de"
ANSCHRIFT = "Straße am Flugplatz 6a, 12487 Berlin"
KONTAKT = "info@ing-bassam.de · +49 176 23581339"

# Farben der Website (fachwissen/artikel.css)
INK = (0x0a, 0x12, 0x1d)
INK_TEXT = (0x0f, 0x1a, 0x28)
TEXT_2 = (0x47, 0x52, 0x62)
LINIE = (0xc9, 0xc5, 0xbb)
AKZENT = (0xf2, 0x6b, 0x2a)
AKZENT_DUNKEL = (0xc2, 0x41, 0x0c)

TYPEN = {"Checkliste", "Protokoll", "Musterschreiben", "Tabelle"}
FORMATE = {"pdf", "pdf-ausfuellbar", "docx", "xlsx"}
FORMAT_NAMEN = {"pdf": "PDF zum Ausdrucken", "pdf-ausfuellbar": "PDF ausfüllbar",
                "docx": "Word", "xlsx": "Excel"}
PLATZHALTER = re.compile(r"\[[^\]\n]{1,80}\]")


# ---------------------------------------------------------------------------
# Beschreibung lesen und prüfen
# ---------------------------------------------------------------------------

def beschreibung_laden(pfad: Path) -> dict:
    daten = yaml.safe_load(pfad.read_text(encoding="utf-8")) or {}
    fehler = beschreibung_pruefen(daten)
    if fehler:
        for f in fehler:
            print(f"FEHLER: {pfad.name}: {f}")
        raise SystemExit(2)
    return daten


def beschreibung_pruefen(d: dict) -> list[str]:
    fehler = []
    for feld in ("titel", "kurzform", "typ", "stand", "dateien", "einleitung"):
        if not d.get(feld):
            fehler.append(f"Feld „{feld}“ fehlt")
    if d.get("typ") not in TYPEN:
        fehler.append(f"typ muss eines von {sorted(TYPEN)} sein")
    if not re.fullmatch(r"[a-z0-9-]{3,60}", str(d.get("kurzform", ""))):
        fehler.append("kurzform: nur Kleinbuchstaben, Ziffern, Bindestriche")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(d.get("stand", ""))):
        fehler.append("stand muss JJJJ-MM-TT sein")
    unbekannt = set(d.get("dateien") or []) - FORMATE
    if unbekannt:
        fehler.append(f"dateien: unbekannte Formate {sorted(unbekannt)}; erlaubt {sorted(FORMATE)}")
    typ = d.get("typ")
    if typ in ("Checkliste", "Protokoll"):
        if not d.get("abschnitte"):
            fehler.append("abschnitte fehlen")
        for i, a in enumerate(d.get("abschnitte") or [], 1):
            if not a.get("titel") or not a.get("punkte"):
                fehler.append(f"abschnitte[{i}] braucht titel und punkte")
            for j, p in enumerate(a.get("punkte") or [], 1):
                if not (isinstance(p, dict) and p.get("text")):
                    fehler.append(f"abschnitte[{i}].punkte[{j}] braucht ein Feld text")
                elif len(p["text"]) > 220:
                    fehler.append(f"abschnitte[{i}].punkte[{j}]: text länger als 220 Zeichen")
    if typ == "Tabelle" or d.get("tabelle"):
        t = d.get("tabelle") or {}
        if not t.get("spalten"):
            fehler.append("tabelle.spalten fehlen")
        for z in t.get("zeilen") or []:
            if len(z) != len(t.get("spalten") or []):
                fehler.append("tabelle.zeilen: jede Zeile braucht so viele Einträge wie Spalten")
                break
        if "xlsx" in (d.get("dateien") or []) and typ != "Tabelle" and not d.get("tabelle"):
            fehler.append("xlsx nur mit einer tabelle")
    if typ == "Musterschreiben":
        b = d.get("brief") or {}
        for feld in ("betreff", "anrede", "absaetze", "gruss"):
            if not b.get(feld):
                fehler.append(f"brief.{feld} fehlt")
        if "xlsx" in (d.get("dateien") or []):
            fehler.append("Musterschreiben nicht als xlsx")
    if typ != "Tabelle" and "pdf-ausfuellbar" in (d.get("dateien") or []) and typ == "Musterschreiben":
        fehler.append("Musterschreiben nicht als pdf-ausfuellbar – dafür docx")
    return fehler


# ---------------------------------------------------------------------------
# Logo als Bild (für Word und Excel) und als Zeichnung (für PDF)
# ---------------------------------------------------------------------------

EIGENES_LOGO = WURZEL / "assets" / "logo.png"


def logo_png(ziel: Path, hoehe: int = 240) -> Path:
    """Wort-Bild-Marke wie im Kopf der Website: dunkles Quadrat mit „BIB“
    und orangefarbener Linie, daneben der Name des Büros.

    Liegt unter assets/logo.png ein eigenes Logo (Querformat, transparenter
    Hintergrund), wird es statt der gezeichneten Marke verwendet.
    """
    if EIGENES_LOGO.exists():
        return EIGENES_LOGO
    from PIL import Image, ImageDraw, ImageFont
    faktor = hoehe / 64
    breite = int(hoehe * 5.2)
    bild = Image.new("RGBA", (breite, hoehe), (0, 0, 0, 0))
    z = ImageDraw.Draw(bild)
    z.rounded_rectangle((0, 0, hoehe - 1, hoehe - 1), radius=int(12 * faktor), fill=INK + (255,))
    z.line((16 * faktor, 46 * faktor, 48 * faktor, 46 * faktor), fill=AKZENT + (255,),
           width=max(2, int(3 * faktor)))
    fira = ImageFont.truetype(str(SCHRIFTEN / "FiraSans-300.ttf"), int(24 * faktor))
    text = "BIB"
    box = z.textbbox((0, 0), text, font=fira)
    z.text(((hoehe - (box[2] - box[0])) / 2 - box[0], 38 * faktor - box[3]), text,
           font=fira, fill=(255, 255, 255, 255))
    gross = ImageFont.truetype(str(SCHRIFTEN / "FiraSans-300.ttf"), int(30 * faktor))
    klein = ImageFont.truetype(str(SCHRIFTEN / "Inter-600.ttf"), int(11 * faktor))
    x = hoehe + int(18 * faktor)
    z.text((x, int(10 * faktor)), "BIB", font=gross, fill=INK_TEXT + (255,))
    z.text((x, int(44 * faktor)), "INGENIEURBÜRO FÜR BAUWESEN", font=klein,
           fill=TEXT_2 + (255,), spacing=0)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    bild.save(ziel)
    return ziel


def logo_pdf(canv, x: float, y: float, hoehe: float = 34) -> None:
    """Dasselbe Logo als Vektorzeichnung auf einem reportlab-Canvas; x, y = linke untere Ecke."""
    from reportlab.lib.colors import Color
    if EIGENES_LOGO.exists():
        from PIL import Image
        with Image.open(EIGENES_LOGO) as bild:
            verhaeltnis = bild.width / bild.height
        canv.drawImage(str(EIGENES_LOGO), x, y, width=hoehe * verhaeltnis, height=hoehe,
                       mask="auto", preserveAspectRatio=True)
        return
    f = hoehe / 64
    canv.saveState()
    canv.setFillColor(Color(*[c / 255 for c in INK]))
    canv.roundRect(x, y, hoehe, hoehe, 12 * f, stroke=0, fill=1)
    canv.setStrokeColor(Color(*[c / 255 for c in AKZENT]))
    canv.setLineWidth(3 * f)
    canv.line(x + 16 * f, y + 18 * f, x + 48 * f, y + 18 * f)
    canv.setFillColor(Color(1, 1, 1))
    canv.setFont("FiraSans-Light", 24 * f)
    canv.drawCentredString(x + hoehe / 2, y + 27 * f, "BIB")
    tx = x + hoehe + 12 * f
    canv.setFillColor(Color(*[c / 255 for c in INK_TEXT]))
    canv.setFont("FiraSans-Light", 26 * f)
    canv.drawString(tx, y + hoehe - 30 * f, "BIB")
    canv.setFillColor(Color(*[c / 255 for c in TEXT_2]))
    canv.setFont("Inter-SemiBold", 9 * f)
    canv.drawString(tx, y + 6 * f, "INGENIEURBÜRO FÜR BAUWESEN")
    canv.restoreState()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _schriften_registrieren() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for name, datei in (("Inter", "Inter-400.ttf"), ("Inter-SemiBold", "Inter-600.ttf"),
                        ("SpaceGrotesk", "SpaceGrotesk-600.ttf"), ("FiraSans-Light", "FiraSans-300.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(SCHRIFTEN / datei)))


def pdf_bauen(d: dict, ziel: Path, ausfuellbar: bool) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Flowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    _schriften_registrieren()
    farbe = lambda rgb: colors.Color(*[c / 255 for c in rgb])  # noqa: E731
    quer = d.get("typ") == "Tabelle" and len((d.get("tabelle") or {}).get("spalten") or []) > 4
    seite = landscape(A4) if quer else A4
    rand = 18 * mm
    stand = datum_deutsch(str(d["stand"]))
    fuss = f"{FIRMA_KURZ} · {WEBSITE} · Stand {stand}"

    st_titel = ParagraphStyle("t", fontName="SpaceGrotesk", fontSize=19, leading=23,
                              textColor=farbe(INK_TEXT), spaceAfter=2)
    st_unter = ParagraphStyle("u", fontName="Inter", fontSize=10.5, leading=14, textColor=farbe(TEXT_2),
                              spaceAfter=8)
    st_text = ParagraphStyle("p", fontName="Inter", fontSize=9.5, leading=13.5, textColor=farbe(INK_TEXT),
                             alignment=TA_LEFT)
    st_hinweis = ParagraphStyle("h", parent=st_text, fontSize=8.2, leading=11, textColor=farbe(TEXT_2))
    st_klein = ParagraphStyle("k", parent=st_text, fontSize=7.8, leading=10.5, textColor=farbe(TEXT_2))
    st_h2 = ParagraphStyle("h2", fontName="SpaceGrotesk", fontSize=12, leading=15,
                           textColor=farbe(INK_TEXT), spaceBefore=10, spaceAfter=4)

    breite = seite[0] - 2 * rand
    feld_nr = [0]

    class Eingabe(Flowable):
        """Beschriftetes Eingabefeld (Kopffelder): Linie oder Formularfeld."""

        def __init__(self, beschriftung: str, hoehe: float = 24):
            super().__init__()
            self.beschriftung, self.h = beschriftung, hoehe

        def wrap(self, aw, ah):
            self.w = aw
            return aw, self.h

        def draw(self):
            c = self.canv
            c.setFont("Inter", 8.5)
            c.setFillColor(farbe(TEXT_2))
            c.drawString(0, self.h - 9, self.beschriftung)
            if ausfuellbar:
                feld_nr[0] += 1
                ax, ay = c.absolutePosition(0, 0)
                c.acroForm.textfield(name=f"feld{feld_nr[0]}", x=ax, y=ay, width=self.w, height=13,
                                     borderWidth=0.5, borderColor=farbe(LINIE),
                                     fillColor=colors.Color(0.98, 0.98, 0.97), fontName="Helvetica",
                                     fontSize=9, textColor=farbe(INK_TEXT), forceBorder=True)
            else:
                c.setStrokeColor(farbe(LINIE))
                c.setLineWidth(0.6)
                c.line(0, 1, self.w, 1)

    class Punkt(Flowable):
        """Eine Checklisten-Zeile: Kästchen, Text, optionaler Hinweis, Bemerkungsfeld."""

        def __init__(self, text: str, hinweis: str, mit_bemerkung: bool):
            super().__init__()
            self.text, self.hinweis, self.mit_bemerkung = text, hinweis, mit_bemerkung
            self.absatz = Paragraph(text, st_text)
            self.hinweis_absatz = Paragraph(hinweis, st_hinweis) if hinweis else None

        def wrap(self, aw, ah):
            self.w = aw
            self.textbreite = aw - 22 - (0.32 * aw if self.mit_bemerkung else 0)
            _, h1 = self.absatz.wrap(self.textbreite, ah)
            h2 = 0
            if self.hinweis_absatz:
                _, h2 = self.hinweis_absatz.wrap(self.textbreite, ah)
                h2 += 1
            self.h = max(h1 + h2 + 7, 18)
            self.h1, self.h2 = h1, h2
            return aw, self.h

        def draw(self):
            c = self.canv
            oben = self.h - 3
            if ausfuellbar:
                feld_nr[0] += 1
                ax, ay = c.absolutePosition(1, oben - 11)
                c.acroForm.checkbox(name=f"k{feld_nr[0]}", x=ax, y=ay, size=10,
                                    borderWidth=0.6, borderColor=farbe(TEXT_2),
                                    fillColor=colors.white, buttonStyle="check", forceBorder=True)
            else:
                c.setStrokeColor(farbe(TEXT_2))
                c.setLineWidth(0.6)
                c.rect(1, oben - 11, 10, 10, stroke=1, fill=0)
            self.absatz.drawOn(c, 22, self.h - self.h1 - 3)
            if self.hinweis_absatz:
                self.hinweis_absatz.drawOn(c, 22, self.h - self.h1 - self.h2 - 3)
            if self.mit_bemerkung:
                bx = 22 + self.textbreite + 8
                bw = self.w - bx
                if ausfuellbar:
                    feld_nr[0] += 1
                    ax, ay = c.absolutePosition(bx, oben - 13)
                    c.acroForm.textfield(name=f"b{feld_nr[0]}", x=ax, y=ay, width=bw, height=13,
                                         borderWidth=0.5, borderColor=farbe(LINIE),
                                         fillColor=colors.Color(0.98, 0.98, 0.97), fontName="Helvetica",
                                         fontSize=8, forceBorder=True)
                else:
                    c.setStrokeColor(farbe(LINIE))
                    c.setLineWidth(0.5)
                    c.line(bx, oben - 12, bx + bw, oben - 12)
            c.setStrokeColor(colors.Color(0.9, 0.89, 0.86))
            c.setLineWidth(0.4)
            c.line(0, 0, self.w, 0)

    def seite_zeichnen(canv, doc):
        canv.saveState()
        # Wasserzeichen: dezent, diagonal, hinter dem Text
        canv.setFillColor(colors.Color(0.06, 0.1, 0.16, alpha=0.06))
        canv.setFont("SpaceGrotesk", 60 if not quer else 70)
        canv.translate(seite[0] / 2, seite[1] / 2)
        canv.rotate(35)
        canv.drawCentredString(0, 0, WEBSITE)
        canv.restoreState()
        canv.saveState()
        logo_pdf(canv, rand, seite[1] - rand - 22, 30)
        canv.setFont("Inter", 8)
        canv.setFillColor(farbe(TEXT_2))
        canv.drawRightString(seite[0] - rand, seite[1] - rand - 4, f"Stand {stand}")
        canv.drawRightString(seite[0] - rand, seite[1] - rand - 15, WEBSITE)
        canv.setStrokeColor(farbe(LINIE))
        canv.setLineWidth(0.6)
        canv.line(rand, seite[1] - rand - 30, seite[0] - rand, seite[1] - rand - 30)
        canv.line(rand, rand + 14, seite[0] - rand, rand + 14)
        canv.setFont("Inter", 7.5)
        canv.drawString(rand, rand + 4, fuss)
        canv.drawRightString(seite[0] - rand, rand + 4, f"Seite {doc.page}")
        canv.restoreState()

    doc = SimpleDocTemplate(str(ziel), pagesize=seite, leftMargin=rand, rightMargin=rand,
                            topMargin=rand + 38, bottomMargin=rand + 22,
                            title=d["titel"], author=FIRMA_LANG, subject=d.get("untertitel", ""),
                            creator=f"{FIRMA_KURZ} – {WEBSITE}")
    teile = [Paragraph(d["titel"], st_titel)]
    if d.get("untertitel"):
        teile.append(Paragraph(d["untertitel"], st_unter))
    teile.append(Paragraph(d["einleitung"], st_text))
    teile.append(Spacer(1, 8))

    kopffelder = d.get("kopffelder") or []
    if kopffelder and d["typ"] != "Musterschreiben":
        zeilen, reihe = [], []
        for feld in kopffelder:
            reihe.append(Eingabe(feld))
            if len(reihe) == 2:
                zeilen.append(reihe)
                reihe = []
        if reihe:
            reihe.append(Spacer(1, 1))
            zeilen.append(reihe)
        tabelle = Table(zeilen, colWidths=[breite / 2 - 6, breite / 2 - 6], hAlign="LEFT")
        tabelle.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                     ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        teile += [tabelle, Spacer(1, 6)]

    typ = d["typ"]
    if typ in ("Checkliste", "Protokoll"):
        mit_bemerkung = typ == "Protokoll"
        for abschnitt in d["abschnitte"]:
            block = [Paragraph(abschnitt["titel"], st_h2)]
            if mit_bemerkung:
                block.append(Paragraph("<font size=7.5 color='#475262'>Bemerkung / Befund</font>",
                                       ParagraphStyle("r", parent=st_klein, alignment=2)))
            for p in abschnitt["punkte"]:
                block.append(Punkt(p["text"], p.get("hinweis", ""), mit_bemerkung))
            teile.append(KeepTogether(block[:3]))
            teile += block[3:]
        if d.get("unterschriften"):
            teile.append(Spacer(1, 22))
            reihe = [Eingabe(f"Ort, Datum, Unterschrift {name}", 30) for name in d["unterschriften"]]
            tabelle = Table([reihe], colWidths=[breite / len(reihe) - 6] * len(reihe), hAlign="LEFT")
            tabelle.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                                         ("RIGHTPADDING", (0, 0), (-1, -1), 12)]))
            teile.append(KeepTogether([tabelle]))

    if typ == "Musterschreiben":
        b = d["brief"]
        adresse = [Paragraph("[Absender: Name, Firma, Anschrift]", st_text), Spacer(1, 14),
                   Paragraph("[Empfänger: Name, Firma, Anschrift]", st_text), Spacer(1, 14),
                   Paragraph("[Ort], [Datum]", ParagraphStyle("r", parent=st_text, alignment=2)),
                   Spacer(1, 12),
                   Paragraph(f"<b>{b['betreff']}</b>", ParagraphStyle("b", parent=st_text,
                                                                     fontName="Inter-SemiBold")),
                   Spacer(1, 10), Paragraph(b["anrede"], st_text), Spacer(1, 6)]
        teile += adresse
        for absatz in b["absaetze"]:
            teile += [Paragraph(absatz, st_text), Spacer(1, 6)]
        teile += [Spacer(1, 8), Paragraph(b["gruss"], st_text), Spacer(1, 26),
                  Paragraph("[Name, Funktion, Unterschrift]", st_text)]
        if b.get("anlagen"):
            teile += [Spacer(1, 10), Paragraph("Anlagen: " + "; ".join(b["anlagen"]), st_hinweis)]

    if d.get("tabelle"):
        t = d["tabelle"]
        spalten = t["spalten"]
        rel = t.get("breiten") or [1] * len(spalten)
        summe = sum(rel)
        cw = [breite * r / summe for r in rel]
        kopf = [Paragraph(f"<b>{s}</b>", ParagraphStyle("th", parent=st_klein, fontName="Inter-SemiBold",
                                                          textColor=farbe(INK_TEXT))) for s in spalten]
        daten = [kopf]
        for zeile in t.get("zeilen") or []:
            daten.append([Paragraph(str(z), st_klein) for z in zeile])
        for _ in range(int(t.get("leerzeilen") or (12 if not t.get("zeilen") else 4))):
            daten.append([""] * len(spalten))
        if typ == "Tabelle":
            teile.append(Spacer(1, 4))
        else:
            teile.append(Paragraph(t.get("titel", "Tabelle"), st_h2))
        tabelle = Table(daten, colWidths=cw, repeatRows=1, hAlign="LEFT")
        tabelle.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, farbe(LINIE)),
            ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.96, 0.955, 0.94)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.985, 0.98, 0.97)]),
            ("MINROWHEIGHT", (0, 1), (-1, -1), 20),
            ("FONTNAME", (0, 1), (-1, -1), "Inter"), ("FONTSIZE", (0, 1), (-1, -1), 8),
        ]))
        teile.append(tabelle)

    if d.get("hinweise"):
        teile.append(Spacer(1, 14))
        teile.append(Paragraph("Hinweise", st_h2))
        for h in d["hinweise"]:
            teile.append(Paragraph("• " + h, st_hinweis))
    if d.get("quellen"):
        teile.append(Spacer(1, 8))
        teile.append(Paragraph("Grundlagen: " + " · ".join(d["quellen"]), st_klein))
    teile.append(Spacer(1, 10))
    teile.append(Paragraph(
        f"Diese Vorlage stellt {FIRMA_LANG} kostenlos bereit ({WEBSITE}). Sie ersetzt keine "
        "Begutachtung des Einzelfalls und keine Rechtsberatung. Weitergabe unverändert und mit "
        "diesem Hinweis erlaubt.", st_klein))

    doc.build(teile, onFirstPage=seite_zeichnen, onLaterPages=seite_zeichnen)
    return ziel


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def docx_bauen(d: dict, ziel: Path, logo: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    stand = datum_deutsch(str(d["stand"]))
    typ = d["typ"]
    quer = typ == "Tabelle" and len((d.get("tabelle") or {}).get("spalten") or []) > 4

    sec = doc.sections[0]
    if quer:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for seite in ("left_margin", "right_margin"):
        setattr(sec, seite, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(*INK_TEXT)

    # Kopfzeile: Logo links, Stand rechts
    kopf = sec.header.paragraphs[0]
    kopf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kopf.add_run().add_picture(str(logo), height=Cm(1.05))
    kopf.add_run("\t\t")
    lauf = kopf.add_run(f"Stand {stand} · {WEBSITE}")
    lauf.font.size = Pt(8)
    lauf.font.color.rgb = RGBColor(*TEXT_2)
    # Wasserzeichen: dezenter Schriftzug als WordArt in der Kopfzeile (Word zeigt ihn auf jeder Seite)
    kopf._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict>'
        '<v:shape id="Wasserzeichen" o:spid="_x0000_s2049" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;width:460pt;height:110pt;rotation:325;'
        'z-index:-251654144;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin" '
        'o:allowincell="f" fillcolor="#0f1a28" stroked="f">'
        '<v:fill opacity="9830f"/>'
        f'<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{WEBSITE}"/>'
        '</v:shape></w:pict></w:r>'))

    fuss = sec.footer.paragraphs[0]
    fuss.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lauf = fuss.add_run(f"{FIRMA_KURZ} · {WEBSITE} · Stand {stand}")
    lauf.font.size = Pt(7.5)
    lauf.font.color.rgb = RGBColor(*TEXT_2)

    def ueberschrift(text: str, groesse: float, farbe=INK_TEXT, vor: float = 8, nach: float = 3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(vor), Pt(nach)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(groesse)
        r.font.color.rgb = RGBColor(*farbe)
        return p

    def absatz(text: str, groesse: float = 10.5, farbe=INK_TEXT, nach: float = 4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(nach)
        r = p.add_run(text)
        r.font.size = Pt(groesse)
        r.font.color.rgb = RGBColor(*farbe)
        return p

    ueberschrift(d["titel"], 17, vor=0, nach=1)
    if d.get("untertitel"):
        absatz(d["untertitel"], 11, TEXT_2, 6)
    absatz(d["einleitung"], nach=8)

    kopffelder = d.get("kopffelder") or []
    if kopffelder and typ != "Musterschreiben":
        t = doc.add_table(rows=(len(kopffelder) + 1) // 2, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, feld in enumerate(kopffelder):
            zelle = t.cell(i // 2, i % 2)
            zelle.text = ""
            p = zelle.paragraphs[0]
            r = p.add_run(f"{feld}: ")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(*TEXT_2)
            p.add_run("_" * 26).font.color.rgb = RGBColor(*LINIE)
        doc.add_paragraph()

    if typ in ("Checkliste", "Protokoll"):
        for a in d["abschnitte"]:
            ueberschrift(a["titel"], 12)
            for p in a["punkte"]:
                z = doc.add_paragraph()
                z.paragraph_format.space_after = Pt(2)
                z.paragraph_format.left_indent = Cm(0.7)
                z.paragraph_format.first_line_indent = Cm(-0.7)
                z.add_run("☐  ").font.size = Pt(12)
                z.add_run(p["text"])
                if p.get("hinweis"):
                    r = z.add_run("  – " + p["hinweis"])
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(*TEXT_2)
                if typ == "Protokoll":
                    b = doc.add_paragraph()
                    b.paragraph_format.left_indent = Cm(0.7)
                    b.paragraph_format.space_after = Pt(6)
                    r = b.add_run("Bemerkung: " + "_" * 60)
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(*LINIE)
        if d.get("unterschriften"):
            doc.add_paragraph()
            t = doc.add_table(rows=2, cols=len(d["unterschriften"]))
            for i, name in enumerate(d["unterschriften"]):
                t.cell(0, i).text = "\n\n"
                t.cell(1, i).text = f"Ort, Datum, Unterschrift {name}"
                for p in t.cell(1, i).paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(*TEXT_2)

    if typ == "Musterschreiben":
        b = d["brief"]
        absatz("[Absender: Name, Firma, Anschrift]", nach=12)
        absatz("[Empfänger: Name, Firma, Anschrift]", nach=12)
        p = absatz("[Ort], [Datum]", nach=12)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ueberschrift(b["betreff"], 11, vor=4, nach=8)
        absatz(b["anrede"], nach=6)
        for text in b["absaetze"]:
            absatz(text, nach=6)
        absatz(b["gruss"], nach=28)
        absatz("[Name, Funktion, Unterschrift]")
        if b.get("anlagen"):
            absatz("Anlagen: " + "; ".join(b["anlagen"]), 9, TEXT_2)

    if d.get("tabelle"):
        t = d["tabelle"]
        if typ != "Tabelle":
            ueberschrift(t.get("titel", "Tabelle"), 12)
        zeilen = list(t.get("zeilen") or [])
        leer = int(t.get("leerzeilen") or (12 if not zeilen else 4))
        tab = doc.add_table(rows=1 + len(zeilen) + leer, cols=len(t["spalten"]))
        tab.style = "Table Grid"
        for i, s in enumerate(t["spalten"]):
            zelle = tab.cell(0, i)
            zelle.text = ""
            r = zelle.paragraphs[0].add_run(str(s))
            r.bold = True
            r.font.size = Pt(9)
            zelle._tc.get_or_add_tcPr().append(parse_xml(
                f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F4F3EF"/>'))
        for zi, zeile in enumerate(zeilen, 1):
            for si, wert in enumerate(zeile):
                tab.cell(zi, si).text = str(wert)
                for p in tab.cell(zi, si).paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

    if d.get("hinweise"):
        ueberschrift("Hinweise", 11, vor=12)
        for h in d["hinweise"]:
            absatz("• " + h, 9, TEXT_2, 2)
    if d.get("quellen"):
        absatz("Grundlagen: " + " · ".join(d["quellen"]), 8, TEXT_2, 8)
    absatz(f"Diese Vorlage stellt {FIRMA_LANG} kostenlos bereit ({WEBSITE}). Sie ersetzt keine "
           "Begutachtung des Einzelfalls und keine Rechtsberatung. Weitergabe unverändert und mit "
           "diesem Hinweis erlaubt.", 8, TEXT_2)

    doc.core_properties.title = d["titel"]
    doc.core_properties.author = FIRMA_LANG
    doc.core_properties.comments = f"{WEBSITE} – Stand {stand}"
    doc.save(str(ziel))
    return ziel


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def xlsx_bauen(d: dict, ziel: Path, logo: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    t = d["tabelle"]
    spalten = t["spalten"]
    stand = datum_deutsch(str(d["stand"]))
    wb = Workbook()
    ws = wb.active
    ws.title = (d["titel"][:28] or "Vorlage").replace("/", "-")

    bild = XlImage(str(logo))
    bild.height, bild.width = 40, int(40 * 5.2)
    ws.add_image(bild, "A1")
    ws.row_dimensions[1].height = 34
    ws["A3"] = d["titel"]
    ws["A3"].font = Font(name="Calibri", size=15, bold=True, color="0F1A28")
    ws["A4"] = d.get("untertitel") or d["einleitung"]
    ws["A4"].font = Font(name="Calibri", size=10, color="475262")
    ws["A5"] = f"Stand {stand} · {FIRMA_KURZ} · {WEBSITE}"
    ws["A5"].font = Font(name="Calibri", size=9, color="475262")

    kopfzeile = 7
    duenn = Side(style="thin", color="C9C5BB")
    rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
    for i, s in enumerate(spalten, 1):
        z = ws.cell(row=kopfzeile, column=i, value=str(s))
        z.font = Font(name="Calibri", bold=True, size=10, color="0F1A28")
        z.fill = PatternFill("solid", fgColor="F4F3EF")
        z.border = rahmen
        z.alignment = Alignment(vertical="center", wrap_text=True)
    rel = t.get("breiten") or [1] * len(spalten)
    for i, r in enumerate(rel, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, 12 * r / max(rel) * 2.2)
    zeilen = list(t.get("zeilen") or [])
    leer = int(t.get("leerzeilen") or (30 if not zeilen else 10))
    for zi in range(len(zeilen) + leer):
        for si in range(len(spalten)):
            wert = zeilen[zi][si] if zi < len(zeilen) else None
            z = ws.cell(row=kopfzeile + 1 + zi, column=si + 1, value=wert)
            z.border = rahmen
            z.alignment = Alignment(vertical="top", wrap_text=True)
            z.font = Font(name="Calibri", size=10)
    ws.freeze_panes = ws.cell(row=kopfzeile + 1, column=1)
    ws.print_title_rows = f"{kopfzeile}:{kopfzeile}"
    ws.page_setup.orientation = "landscape" if len(spalten) > 4 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = f"{FIRMA_KURZ} · {WEBSITE} · Stand {stand}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Seite &P von &N"
    ws.oddFooter.right.size = 8
    ws.oddHeader.center.text = f"&8&K475262{WEBSITE}"

    if d.get("hinweise") or d.get("quellen"):
        hw = wb.create_sheet("Hinweise")
        hw["A1"] = d["titel"]
        hw["A1"].font = Font(bold=True, size=12)
        hw["A2"] = d["einleitung"]
        r = 4
        for h in d.get("hinweise") or []:
            hw.cell(row=r, column=1, value="• " + h)
            r += 1
        if d.get("quellen"):
            r += 1
            hw.cell(row=r, column=1, value="Grundlagen: " + " · ".join(d["quellen"]))
        r += 2
        hw.cell(row=r, column=1, value=f"Diese Vorlage stellt {FIRMA_LANG} kostenlos bereit ({WEBSITE}). "
                                       "Sie ersetzt keine Begutachtung des Einzelfalls und keine Rechtsberatung.")
        hw.column_dimensions["A"].width = 120
    wb.properties.title = d["titel"]
    wb.properties.creator = FIRMA_LANG
    wb.save(str(ziel))
    return ziel


# ---------------------------------------------------------------------------
# Ablauf
# ---------------------------------------------------------------------------

def datum_deutsch(iso: str) -> str:
    j, m, t = iso.split("-")
    return f"{t}.{m}.{j}"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("beschreibung", type=Path)
    parser.add_argument("--ziel", type=Path, default=None, help="Ordner für die Dateien")
    parser.add_argument("--pruefen", action="store_true", help="nur die Beschreibung prüfen")
    args = parser.parse_args()

    d = beschreibung_laden(args.beschreibung)
    if args.pruefen:
        print(f"{args.beschreibung.name}: Beschreibung in Ordnung ({d['typ']}, "
              f"{', '.join(FORMAT_NAMEN[f] for f in d['dateien'])}).")
        return 0

    ziel = args.ziel or (WURZEL / "vorlagen" / d["kurzform"])
    ziel.mkdir(parents=True, exist_ok=True)
    basis = d["kurzform"]
    # Das Logo-Bild braucht nur Word und Excel; es entsteht im Temp-Ordner und
    # wird nicht ins Repository geschrieben.
    logo = logo_png(Path(tempfile.gettempdir()) / "bib-logo.png")
    erzeugt: list[Path] = []
    for fmt in d["dateien"]:
        if fmt == "pdf":
            erzeugt.append(pdf_bauen(d, ziel / f"{basis}.pdf", ausfuellbar=False))
        elif fmt == "pdf-ausfuellbar":
            erzeugt.append(pdf_bauen(d, ziel / f"{basis}-ausfuellbar.pdf", ausfuellbar=True))
        elif fmt == "docx":
            erzeugt.append(docx_bauen(d, ziel / f"{basis}.docx", logo))
        elif fmt == "xlsx":
            erzeugt.append(xlsx_bauen(d, ziel / f"{basis}.xlsx", logo))
    for p in erzeugt:
        print(p.relative_to(WURZEL).as_posix() if WURZEL in p.resolve().parents else p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
