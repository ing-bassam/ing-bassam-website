#!/usr/bin/env python3
"""Wandelt daten/marktdaten.xlsx in wertrechner/data/marktdaten.json um.

Die Excel-Datei ist die Pflegedatei des Büros: Dort werden neue Werte des
Gutachterausschusses eingetragen (Sachwertfaktoren, Liegenschaftszinssätze,
Baupreisindex, Bewirtschaftungskosten …). Der Browser liest nie die Excel-Datei,
sondern nur die hier erzeugte JSON-Datei.

Regeln für die Excel-Datei (siehe daten/ANLEITUNG.md):
  - Jedes Tabellenblatt wird zu einem Eintrag unter "blaetter".
  - Die erste Zeile enthält die Spaltennamen, jede weitere Zeile einen Datensatz.
  - Leere Zeilen werden übersprungen, leere Zellen werden zu null.
  - Blätter, deren Name mit "_" beginnt, werden nicht übernommen (Notizen).
  - Jede Zeile soll eine Spalte "Quelle/Stand" haben; sie bleibt erhalten, damit
    der Rechenweg auf der Website die Quelle nennen kann.

Aufruf:
    python tools/marktdaten_bauen.py            schreibt die JSON-Datei
    python tools/marktdaten_bauen.py --pruefen  nur prüfen, nichts schreiben
Rückgabewert 1, wenn Pflichtblätter oder Pflichtschlüssel fehlen.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.exit("openpyxl fehlt: python -m pip install openpyxl")

WURZEL = Path(__file__).resolve().parent.parent
QUELLE = WURZEL / "daten" / "marktdaten.xlsx"
ZIEL = WURZEL / "wertrechner" / "data" / "marktdaten.json"

# Diese Blätter und Schlüssel braucht der Rechenkern zwingend.
PFLICHTBLAETTER = [
    "Modellparameter", "Sachwertfaktoren", "SWF_Korrekturen", "Liegenschaftszinssaetze",
    "Bewirtschaftungskosten", "RND", "NHK2010", "GFZ_Koeffizienten", "Ortsteile",
    "Gueltigkeit", "Baupreisindex", "Quellen",
]
PFLICHTSCHLUESSEL = {
    "Modellparameter": [
        "stichtag_faktoren", "brw_stichtag_modell", "brw_stichtag_aktuell", "gnd_wohnen",
        "regionalfaktor", "bpi_2010", "spanne_prozent", "swf_rundung_nachkommastellen",
        "lz_rundung_nachkommastellen",
    ],
    "Sachwertfaktoren": ["konstante", "koeff_sachwert", "koeff_tag", "tage", "gruppe_1", "gruppe_2", "gruppe_3"],
    "Liegenschaftszinssaetze": [
        "konstante", "koeff_miete", "koeff_tag", "tage", "zuschlag_city", "zuschlag_ost", "zuschlag_west",
        "gewerbe_je_prozentpunkt", "marktanpassungsfaktor",
    ],
}


def wert(zelle):
    v = zelle.value
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def blatt_lesen(ws) -> list[dict]:
    zeilen = list(ws.iter_rows())
    if not zeilen:
        return []
    spalten = [wert(z) for z in zeilen[0]]
    datensaetze = []
    for zeile in zeilen[1:]:
        werte = [wert(z) for z in zeile]
        if all(v is None for v in werte):
            continue
        datensaetze.append({s: v for s, v in zip(spalten, werte) if s is not None})
    return datensaetze


def schluessel_werte(zeilen: list[dict]) -> dict:
    """Blätter mit den Spalten Schlüssel/Wert werden zusätzlich als Wörterbuch angeboten."""
    return {z["Schlüssel"]: z.get("Wert") for z in zeilen if z.get("Schlüssel")}


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--pruefen", action="store_true")
    a = p.parse_args()

    if not QUELLE.exists():
        print(f"FEHLER: {QUELLE.relative_to(WURZEL)} fehlt.")
        return 1
    wb = load_workbook(QUELLE, data_only=True, read_only=True)
    blaetter = {ws.title: blatt_lesen(ws) for ws in wb.worksheets if not ws.title.startswith("_")}

    fehler = [f"Blatt fehlt: {b}" for b in PFLICHTBLAETTER if b not in blaetter]
    for blatt, schluessel in PFLICHTSCHLUESSEL.items():
        if blatt in blaetter:
            vorhanden = schluessel_werte(blaetter[blatt])
            fehler += [f"{blatt}: Schlüssel fehlt oder leer: {s}" for s in schluessel if vorhanden.get(s) is None]
    for blatt, zeilen in blaetter.items():
        ohne_quelle = [i + 2 for i, z in enumerate(zeilen) if not z.get("Quelle/Stand")]
        if ohne_quelle:
            fehler.append(f"{blatt}: Zeile(n) ohne „Quelle/Stand“: {ohne_quelle[:10]}")
    if fehler:
        print("FEHLER in daten/marktdaten.xlsx:")
        for f in fehler:
            print("  -", f)
        return 1

    ausgabe = {
        "quelle": "daten/marktdaten.xlsx",
        "parameter": schluessel_werte(blaetter["Modellparameter"]),
        "sachwertfaktor": schluessel_werte(blaetter["Sachwertfaktoren"]),
        "liegenschaftszins": schluessel_werte(blaetter["Liegenschaftszinssaetze"]),
        "blaetter": blaetter,
    }
    text = json.dumps(ausgabe, ensure_ascii=False, indent=1) + "\n"
    if a.pruefen:
        print(f"OK: {len(blaetter)} Blätter, {sum(len(z) for z in blaetter.values())} Zeilen, {len(text)} Zeichen JSON.")
        return 0
    ZIEL.parent.mkdir(parents=True, exist_ok=True)
    alt = ZIEL.read_text(encoding="utf-8") if ZIEL.exists() else None
    if alt == text:
        print(f"Unverändert: {ZIEL.relative_to(WURZEL)}")
    else:
        ZIEL.write_text(text, encoding="utf-8")
        print(f"Geschrieben: {ZIEL.relative_to(WURZEL)} ({len(text) / 1e3:.0f} kB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
